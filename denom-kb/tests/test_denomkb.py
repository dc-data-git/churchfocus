"""Offline tests: a miniature workbook with the real structure, mocked web, fake model."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest
import yaml

from denomkb import pipeline as pl
from denomkb.cli import main
from denomkb.index import BM25, chunk_docs, quote_found
from denomkb.workbook import load

FIELDS = [  # (field_id, layer, sheet label)
    ("identity.primary_family", "identity", "Primary family"),
    ("identity.official_website", "identity", "Official website"),
    ("history.founding_date", "history", "Founding date"),
    ("theology.baptism_mode", "theology", "Baptism mode"),
    ("governance.women_ordination", "governance", "Women ordination"),
    ("practice_culture.alcohol", "practice_culture", "Alcohol"),
    ("distinguishing_metadata.distinctives", "distinguishing_metadata", "Distinctives"),
]
SHEETS = {"identity": "1 Identity", "history": "2 History", "theology": "3 Theology", "governance": "4 Governance",
          "practice_culture": "5 Practice and Culture", "distinguishing_metadata": "6 Distinguishing Metadata"}
GROUPS = [("Test Baptist Fellowship", "tbf", "900", 120, 50000),
          ("Test Jewish Movement", "tjm", "901", 40, 9000)]
EXISTING = {("tbf", "identity.primary_family"): "Baptist", ("tbf", "theology.baptism_mode"): "Immersion"}


def build_workbook(path: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(name, title, header, rows):
        ws = wb.create_sheet(name)
        ws.append([None])
        ws.append([title])
        ws.append([None])
        ws.append([None])
        ws.append(header)
        for r in rows:
            ws.append(r)
    sheet("Groups", "groups", ["Census group name", "Database ID", "Census code", "Canonical database name",
                               "Congregations (2020)", "Adherents (2020)", "Share of total adherents", "Share of US population"],
          [[n, i, c, n, str(k), str(a), "0.001", "0.0001"] for n, i, c, k, a in GROUPS])
    for layer, sh in SHEETS.items():
        labels = [lbl for f, l, lbl in FIELDS if l == layer]
        fids = [f for f, l, lbl in FIELDS if l == layer]
        sheet(sh, sh, ["Census group name", "Database ID", *labels],
              [[n, i, *[EXISTING.get((i, f), "Unknown") for f in fids]] for n, i, *_ in GROUPS])
    sheet("Evidence", "ev", ["Group", "Database ID", "Field ID", "Evidence status", "Source IDs", "Verified date", "Notes"],
          [["Test Baptist Fellowship", "tbf", "theology.baptism_mode", "needs_review", "tbf_src", "Not verified", ""]])
    sheet("Field Dictionary", "fd", ["Field ID", "Layer", "Subtopic", "Description", "Data type", "Definition note"],
          [[f, l, f.split(".")[1], lbl, "string_array_object_or_null", ""] for f, l, lbl in FIELDS])
    sheet("Sources", "src", ["Source ID", "Title", "URL", "Source type", "Retrieval status", "Document date", "Notes"],
          [["tbf_src", "TBF Statement", "https://tbf.example/faith", "official_denomination", "retrieved", "Unknown", ""]])
    wb.save(path)


BAPTIST_PAGE = ("Test Baptist Fellowship statement of faith. Baptism is the immersion of a believer in water in the name "
                "of the Father, the Son, and the Holy Spirit. The fellowship does not ordain women to the office of "
                "pastor, which is limited to men as qualified by Scripture. Members are encouraged to abstain from alcohol "
                "as a matter of witness and personal holiness in every congregation.")
JEWISH_PAGE = ("The Test Jewish Movement is a movement within Judaism in the United States. The movement began to "
               "ordain women as rabbis in 1985 and continues to ordain women today in its seminary programs.")


@pytest.fixture
def project(tmp_path, monkeypatch):
    build_workbook(tmp_path / "kb.xlsx")
    spec = {f: {"description": lbl, "extract": f != "distinguishing_metadata.distinctives",
                "christian_only": f == "theology.baptism_mode",
                "keywords": {"theology.baptism_mode": "baptism immersion sprinkling",
                             "governance.women_ordination": "women ordination ordain women",
                             "practice_culture.alcohol": "alcohol abstain"}.get(f, lbl.lower())}
            for f, _l, lbl in FIELDS}
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "fields.yaml").write_text(yaml.safe_dump(spec))
    cfg = yaml.safe_load((Path(__file__).resolve().parents[1] / "config.example.yaml").read_text())
    cfg["workbook"] = "kb.xlsx"
    (tmp_path / "config.yaml").write_text(yaml.safe_dump(cfg))

    def fake_wikipedia(web, name):
        text = BAPTIST_PAGE if "Baptist" in name else JEWISH_PAGE
        return {"title": name, "pageid": 1, "match_score": 1.0, "url": "https://en.wikipedia.org/wiki/" + name.replace(" ", "_"),
                "text": text, "qid": "Q1" if "Baptist" in name else None}
    monkeypatch.setattr(pl, "wikipedia", fake_wikipedia)
    monkeypatch.setattr(pl, "wikidata", lambda web, qid: {"qid": qid, "url": "https://www.wikidata.org/wiki/Q1",
                                                          "founding_year": "1901"} if qid else {})
    monkeypatch.setattr(pl, "crawl_site", lambda web, start, n, d=2: [])
    return tmp_path


def test_index_and_quotes():
    ch = chunk_docs([{"url": "u", "title": "t", "text": BAPTIST_PAGE, "source_type": "official"}], words=30, overlap=5)
    assert len(ch) >= 2
    top = BM25(ch).search("women ordination ordain", 1)[0]
    assert "ordain women" in top["text"]
    assert quote_found("does not ordain women to the office of pastor", BAPTIST_PAGE)
    assert quote_found("does not ordain women to the office of pastor,", BAPTIST_PAGE)       # punctuation tolerant
    assert not quote_found("the fellowship ordains women as senior pastors", BAPTIST_PAGE)   # fabricated


def test_workbook_load(project):
    kb = load(project / "kb.xlsx")
    assert len(kb.groups) == 2 and len(kb.fields) == len(FIELDS)
    assert kb.values[("tbf", "theology.baptism_mode")] == "Immersion"
    assert kb.source_urls_for("tbf")[0]["url"] == "https://tbf.example/faith"


def test_dry_run_pipeline_and_export(project):
    main(["-c", str(project / "config.yaml"), "run", "--dry-run", "--offline"])
    work = project / "work_dryrun"
    tbf = json.loads((work / "groups" / "tbf" / "proposals.json").read_text())
    tjm = json.loads((work / "groups" / "tjm" / "proposals.json").read_text())
    acts = {(p["field_id"], p["action"]) for p in tbf}
    assert ("history.founding_date", "fill") in acts                      # deterministic Wikidata fill
    assert ("governance.women_ordination", "fill") in acts                # extracted with verified quote
    for p in tbf + tjm:
        if p["action"] == "fill" and p["source_type"] != "wikidata":
            assert quote_found(p["quote"], BAPTIST_PAGE + " " + JEWISH_PAGE)
    assert not any(p["field_id"] == "distinguishing_metadata.distinctives" for p in tbf)   # extract: false
    assert ("theology.baptism_mode", "not_applicable") in {(p["field_id"], p["action"]) for p in tjm}
    assert ("governance.women_ordination", "fill") in {(p["field_id"], p["action"]) for p in tjm}  # not Christian-only
    out = work / "out"
    kbj = json.loads((out / "denominations_kb.json").read_text())
    g = next(x for x in kbj["groups"] if x["id"] == "tbf")
    assert g["fields"]["governance.women_ordination"]["status"] == "model_extracted_needs_review"
    assert g["fields"]["governance.women_ordination"]["evidence"][0]["url"].startswith("https://")
    rows = list(csv.DictReader(open(out / "proposals.csv", encoding="utf-8")))
    assert rows and "review_decision" in rows[0]
    wb = openpyxl.load_workbook(out / "US_Religious_Groups_filled.xlsx")
    ws = wb["4 Governance"]
    cell = ws.cell(row=6, column=3)
    assert cell.value != "Unknown" and cell.comment is not None
    assert "Proposals" in wb.sheetnames
    src = load(project / "kb.xlsx")                      # input untouched
    assert src.values[("tbf", "governance.women_ordination")] == "Unknown"
    assert (out / "report.md").read_text().count("| tbf") == 0      # report lists census names, not ids
    # resumable: second run skips finished groups
    main(["-c", str(project / "config.yaml"), "run", "--dry-run", "--offline"])


def test_fabricated_quote_is_dropped(project):
    from denomkb.cli import _setup
    import argparse
    cfg, root, work, kb, web, llm = _setup(argparse.Namespace(config=str(project / "config.yaml"), dry_run=True, offline=True))
    r = pl.Runner(cfg, root, kb, web, llm, dry_run=True)
    bm = BM25(chunk_docs([{"url": "https://x", "title": "t", "text": BAPTIST_PAGE, "source_type": "official"}]))
    bad = {"fields": [{"field_id": "governance.women_ordination", "status": "stated", "value": "Women ordained",
                       "excerpt": 1, "quote": "the fellowship joyfully ordains women as senior pastors"}]}
    llm.complete_json = lambda *a, **k: bad
    assert r.extract(kb.group("tbf"), bm, ["governance.women_ordination"]) == []


def test_verify_filters_unsupported(project):
    main(["-c", str(project / "config.yaml"), "run", "--dry-run", "--offline"])
    work = project / "work_dryrun"
    path = work / "groups" / "tbf" / "proposals.json"
    props = json.loads(path.read_text())
    # plant a fill whose quote does not support the value (the failure mode the audit found)
    bad = dict(next(p for p in props if p["action"] == "fill" and p["source_type"] != "wikidata"))
    bad.update(field_id="practice_culture.alcohol", new_value="Weekly communion with wine required", quote=BAPTIST_PAGE[:120])
    props.append(bad)
    path.write_text(json.dumps(props))
    main(["-c", str(project / "config.yaml"), "verify", "--dry-run"])
    props = json.loads(path.read_text())
    checked = [p for p in props if p.get("verify")]
    assert checked and all(p["verify"]["version"] == "verify.v1" for p in checked)
    planted = next(p for p in props if p["new_value"] == "Weekly communion with wine required")
    assert planted["verify"]["verdict"] == "unsupported"
    rows = list(csv.DictReader(open(work / "out" / "proposals.csv", encoding="utf-8")))
    r = next(x for x in rows if x["new_value"] == "Weekly communion with wine required")
    assert r["applied"] == "no" and r["verify_verdict"] == "unsupported"
    kbj = json.loads((work / "out" / "denominations_kb.json").read_text())
    g = next(x for x in kbj["groups"] if x["id"] == "tbf")
    assert all(e.get("quote") != BAPTIST_PAGE[:120] or e["action"] != "fill"
               for e in g["fields"]["practice_culture.alcohol"]["evidence"])
    assert "## Verification" in (work / "out" / "report.md").read_text()


def test_same_name():
    assert pl._same_name("Assemblies of God", "Assemblies of God")
    assert pl._same_name("Church of God (Anderson, Indiana)", "Church of God (Anderson, Indiana)")
    assert not pl._same_name("Full Gospel", "Full Gospel Christian Assemblies International")
